# -*- coding: utf-8 -*-
"""
简易 Markdown -> docx（python-docx）。
支持：# 标题, 列表(-/*/数字), 表格(| ... |), 加粗 **x**, 图片 ![alt](path)（本地存在才嵌入）。
用法：python md_to_docx.py <input.md> <output.docx>
"""
import sys, os, re

def add_runs(paragraph, text):
    # 处理 **加粗**
    for i, part in enumerate(re.split(r'(\*\*.+?\*\*)', text)):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        else:
            paragraph.add_run(part)

def main():
    if len(sys.argv) < 3:
        print("usage: md_to_docx.py in.md out.docx"); sys.exit(2)
    src, dst = sys.argv[1], sys.argv[2]
    from docx import Document
    from docx.shared import Pt, Inches
    doc = Document()
    with open(src, encoding='utf-8') as f:
        lines = f.read().split('\n')

    i = 0
    base = os.path.dirname(os.path.abspath(src))
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            lvl = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2).strip(), level=lvl)
            continue
        # 图片 ![alt](path)
        im = re.match(r'^!\[[^\]]*\]\(([^)]+)\)\s*$', line)
        if im:
            p = im.group(1).strip()
            cand = p if os.path.isabs(p) else os.path.join(base, p)
            try:
                if os.path.exists(cand):
                    doc.add_picture(cand, width=Inches(5.5))
                else:
                    doc.add_paragraph(f'[图片缺失: {p}]')
            except Exception:
                doc.add_paragraph(f'[图片: {p}]')
            continue
        # 表格（连续的 | ... | 行）
        if line.lstrip().startswith('|') and '|' in line[1:]:
            rows = []
            j = i - 1
            while j < len(lines) and lines[j].lstrip().startswith('|'):
                cells = [c.strip() for c in lines[j].strip().strip('|').split('|')]
                rows.append(cells); j += 1
            # 去掉分隔行 |---|---|
            rows = [r for r in rows if not all(re.match(r'^:?-{2,}:?$', c or '-') for c in r)]
            i = j
            if rows:
                ncol = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncol)
                table.style = 'Light Grid Accent 1'
                for ri, r in enumerate(rows):
                    cells = table.add_row().cells
                    for ci in range(ncol):
                        cells[ci].text = r[ci] if ci < len(r) else ''
            continue
        # 列表
        lm = re.match(r'^\s*([-*]|\d+[.)])\s+(.*)$', line)
        if lm:
            para = doc.add_paragraph(style='List Bullet' if lm.group(1) in ('-', '*') else 'List Number')
            add_runs(para, lm.group(2))
            continue
        # 普通段落
        para = doc.add_paragraph()
        add_runs(para, line)

    doc.save(dst)
    print("ok:", dst)

if __name__ == '__main__':
    main()
